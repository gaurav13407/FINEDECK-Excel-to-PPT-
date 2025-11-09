"""
Test that mimics EXACTLY what the backend does
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'src'))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pandas as pd
from converter.enhanced_charts import EnhancedChartBuilder

print("=" * 80)
print("BACKEND SIMULATION TEST")
print("=" * 80)

# Create presentation (like backend does)
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Add blank slide (layout 6 like backend uses)
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Add title
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.6))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "Data Insights & Analysis"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = RGBColor(31, 73, 125)

# Create test data (like what backend gets from Excel)
test_data = pd.DataFrame({
    'Product': ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J'],
    'Revenue': [100, 200, 150, 300, 250, 180, 220, 190, 210, 170],
    'Profit': [20, 40, 30, 60, 50, 35, 45, 38, 42, 34],
    'Growth': [5, 10, 7, 15, 12, 8, 11, 9, 10, 8]
})

print("\n📊 Test Data (first 10 rows):")
print(test_data.head(10))

# Initialize EnhancedChartBuilder (like backend does)
template_colors = {
    'navy': (31, 73, 125),
    'light_blue': (68, 114, 196),
    'teal': (91, 155, 213),
    'gray': (165, 165, 165),
    'dark_gray': (89, 89, 89),
    'accent_orange': (237, 125, 49),
    'accent_green': (112, 173, 71)
}

builder = EnhancedChartBuilder(template_colors)

print("\n✨ Using EnhancedChartBuilder with auto-detection...")
print(f"   Calling auto_create_chart with:")
print(f"   - left=Inches(5.2) = {Inches(5.2)} EMUs")
print(f"   - top=Inches(1.3) = {Inches(1.3)} EMUs")
print(f"   - width=Inches(4.3) = {Inches(4.3)} EMUs")
print(f"   - height=Inches(3.5) = {Inches(3.5)} EMUs")

# Create chart EXACTLY like backend does
chart_data = test_data.head(10).copy()

chart = builder.auto_create_chart(
    slide,
    chart_data,
    left=Inches(5.2),  # EXACTLY like backend
    top=Inches(1.3),
    width=Inches(4.3),
    height=Inches(3.5),
    title="Data Insights"
)

if chart:
    print(f"   ✓ Enhanced chart created successfully!")
    print(f"   Chart type: {chart.chart_type}")
else:
    print(f"   ❌ Chart creation returned None!")

# Check slide
print(f"\n📊 Checking slide shapes...")
chart_count = 0
for i, shape in enumerate(slide.shapes):
    print(f"   Shape {i+1}: {shape.shape_type}")
    if hasattr(shape, 'chart'):
        chart_count += 1
        print(f"      ✅ This is a CHART!")
        print(f"      Position: left={shape.left} ({shape.left/914400:.2f}in), top={shape.top} ({shape.top/914400:.2f}in)")
        print(f"      Size: width={shape.width} ({shape.width/914400:.2f}in), height={shape.height} ({shape.height/914400:.2f}in)")

print(f"\nTotal charts on slide: {chart_count}")

# Save
output_path = "backend_simulation_test.pptx"
prs.save(output_path)
print(f"\n💾 Saved to: {output_path}")

# Verify
print("\n" + "=" * 80)
print("VERIFICATION")
print("=" * 80)

prs2 = Presentation(output_path)
slide2 = prs2.slides[0]

print(f"Shapes in saved file: {len(slide2.shapes)}")
chart_count2 = 0

for i, shape in enumerate(slide2.shapes):
    if hasattr(shape, 'chart'):
        chart_count2 += 1
        print(f"\n✅ CHART FOUND!")
        print(f"   Position: left={shape.left} EMUs ({shape.left/914400:.2f} inches)")
        print(f"   Position: top={shape.top} EMUs ({shape.top/914400:.2f} inches)")
        print(f"   Size: width={shape.width} EMUs ({shape.width/914400:.2f} inches)")
        print(f"   Size: height={shape.height} EMUs ({shape.height/914400:.2f} inches)")
        
        expected_left = Inches(5.2)
        expected_top = Inches(1.3)
        
        if shape.left == expected_left and shape.top == expected_top:
            print(f"   ✅ Position MATCHES expected values!")
        else:
            print(f"   ❌ Position WRONG!")
            print(f"      Expected: left={expected_left}, top={expected_top}")
            print(f"      Got: left={shape.left}, top={shape.top}")

if chart_count2 == 0:
    print("\n❌ NO CHARTS in saved file!")
else:
    print(f"\n✅ {chart_count2} chart(s) found in saved file!")
    print("\n👉 Open backend_simulation_test.pptx to see if chart is visible")
