"""
COMPLETE PPT ANALYZER
Analyzes a PowerPoint file to show EVERYTHING about its charts
"""
import sys
from pptx import Presentation
from pptx.util import Inches

if len(sys.argv) < 2:
    print("Usage: python analyze_ppt.py <path_to_pptx_file>")
    print("\nExample:")
    print('  python analyze_ppt.py "downloads/my_presentation.pptx"')
    sys.exit(1)

ppt_file = sys.argv[1]

print("="*80)
print(f"ANALYZING: {ppt_file}")
print("="*80)

try:
    prs = Presentation(ppt_file)
    
    print(f"\n📊 Total Slides: {len(prs.slides)}")
    print(f"📐 Slide Size: {prs.slide_width} x {prs.slide_height} EMUs")
    
    total_charts = 0
    total_shapes = 0
    
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_charts = 0
        slide_shape_count = len(slide.shapes)
        total_shapes += slide_shape_count
        
        print(f"\n{'='*80}")
        print(f"SLIDE {slide_num}")
        print(f"{'='*80}")
        print(f"Total shapes on slide: {slide_shape_count}")
        
        # Check for charts
        for shape_num, shape in enumerate(slide.shapes, 1):
            if hasattr(shape, 'chart'):
                slide_charts += 1
                total_charts += 1
                chart = shape.chart
                
                print(f"\n  📈 CHART FOUND #{slide_charts}")
                print(f"     Shape Index: {shape_num}")
                print(f"     Chart Type: {chart.chart_type} ({chart.chart_type.name if hasattr(chart.chart_type, 'name') else 'Unknown'})")
                
                # Position information
                left_emu = shape.left
                top_emu = shape.top
                width_emu = shape.width
                height_emu = shape.height
                
                left_inches = left_emu / 914400
                top_inches = top_emu / 914400
                width_inches = width_emu / 914400
                height_inches = height_emu / 914400
                
                print(f"\n     POSITION (EMUs):")
                print(f"       Left:   {left_emu:,} EMUs")
                print(f"       Top:    {top_emu:,} EMUs")
                print(f"       Width:  {width_emu:,} EMUs")
                print(f"       Height: {height_emu:,} EMUs")
                
                print(f"\n     POSITION (Inches):")
                print(f"       Left:   {left_inches:.2f} inches")
                print(f"       Top:    {top_inches:.2f} inches")
                print(f"       Width:  {width_inches:.2f} inches")
                print(f"       Height: {height_inches:.2f} inches")
                
                # Check if position is reasonable
                print(f"\n     POSITION CHECK:")
                if left_emu > 10000000000:  # More than 10 billion
                    print(f"       ❌ LEFT is WRONG: {left_emu:,} EMUs = {left_inches:,.2f} inches")
                    print(f"          This is {left_inches/12:,.2f} FEET from left edge!")
                    print(f"          Chart is OFF-SCREEN to the right")
                elif left_emu > 100000000:  # More than 100 million
                    print(f"       ⚠️  LEFT is suspicious: {left_emu:,} EMUs")
                    print(f"          Might be off-screen")
                else:
                    print(f"       ✅ LEFT is OK: {left_inches:.2f} inches")
                
                if top_emu > 10000000000:
                    print(f"       ❌ TOP is WRONG: {top_emu:,} EMUs = {top_inches:,.2f} inches")
                    print(f"          Chart is OFF-SCREEN below")
                elif top_emu > 100000000:
                    print(f"       ⚠️  TOP is suspicious: {top_emu:,} EMUs")
                else:
                    print(f"       ✅ TOP is OK: {top_inches:.2f} inches")
                
                # Slide bounds (typical is 10" x 7.5")
                slide_right = prs.slide_width
                slide_bottom = prs.slide_height
                
                if left_emu + width_emu > slide_right:
                    print(f"       ⚠️  Chart extends PAST right edge of slide")
                if top_emu + height_emu > slide_bottom:
                    print(f"       ⚠️  Chart extends PAST bottom edge of slide")
                
                # Check chart data
                print(f"\n     CHART DATA:")
                try:
                    if hasattr(chart, 'plots') and len(chart.plots) > 0:
                        plot = chart.plots[0]
                        print(f"       Series Count: {len(plot.series)}")
                        
                        for series_num, series in enumerate(plot.series, 1):
                            print(f"       Series {series_num}: '{series.name}' ({len(series.values)} points)")
                            
                            # Check if values are all zero or empty
                            values = list(series.values)
                            if all(v == 0 for v in values):
                                print(f"         ⚠️  All values are ZERO")
                            elif len(values) == 0:
                                print(f"         ⚠️  No data points")
                            else:
                                print(f"         ✅ Has data: min={min(values):.2f}, max={max(values):.2f}")
                    else:
                        print(f"       ⚠️  No plot data found")
                except Exception as e:
                    print(f"       ❌ Error reading chart data: {e}")
                
                # Check if chart has title
                if chart.has_title:
                    title_text = chart.chart_title.text_frame.text
                    print(f"       Title: '{title_text}'")
                else:
                    print(f"       No title")
        
        if slide_charts == 0:
            print(f"  ℹ️  No charts on this slide")
    
    print(f"\n{'='*80}")
    print(f"SUMMARY")
    print(f"{'='*80}")
    print(f"Total Slides: {len(prs.slides)}")
    print(f"Total Shapes: {total_shapes}")
    print(f"Total Charts: {total_charts}")
    
    if total_charts == 0:
        print(f"\n❌ NO CHARTS FOUND IN THIS FILE!")
        print(f"   This means charts are NOT being created at all.")
    else:
        print(f"\n✅ Found {total_charts} chart(s)")
        print(f"   Check the position details above to see if they're visible.")
    
except FileNotFoundError:
    print(f"\n❌ ERROR: File not found: {ppt_file}")
    print(f"   Make sure the path is correct and the file exists.")
except Exception as e:
    print(f"\n❌ ERROR: {e}")
    import traceback
    traceback.print_exc()

print(f"\n{'='*80}")
