"""
Verify improved presentation has better charts and insights
"""

from pptx import Presentation
import os

def verify_improved_presentation():
    """Check the improved presentation"""
    
    ppt_path = "examples/professional_demo/Financials_IMPROVED_AI_PRO.pptx"
    
    print("\n" + "="*80)
    print("🔍 VERIFYING IMPROVED PRESENTATION")
    print("="*80)
    print(f"📁 File: {os.path.basename(ppt_path)}")
    
    if not os.path.exists(ppt_path):
        print(f"❌ File not found!")
        return
    
    # Load presentation
    prs = Presentation(ppt_path)
    file_size_kb = os.path.getsize(ppt_path) / 1024
    
    print(f"📊 Slide Count: {len(prs.slides)}")
    print(f"💾 File Size: {file_size_kb:.1f} KB")
    
    # Analyze each slide
    print(f"\n{'='*80}")
    print("📋 SLIDE-BY-SLIDE ANALYSIS")
    print(f"{'='*80}")
    
    total_charts = 0
    total_tables = 0
    chart_types = []
    
    for slide_idx, slide in enumerate(prs.slides, 1):
        print(f"\n🔹 Slide {slide_idx}:")
        
        # Get slide title
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text and len(text) < 100:
                    title = text
                    break
        
        print(f"   Title: {title[:60]}...")
        
        # Count shapes
        charts_on_slide = 0
        tables_on_slide = 0
        text_boxes = 0
        
        for shape in slide.shapes:
            if shape.has_chart:
                charts_on_slide += 1
                total_charts += 1
                chart_type = str(shape.chart.chart_type)
                chart_types.append(chart_type)
            
            if shape.has_table:
                tables_on_slide += 1
                total_tables += 1
            
            if shape.has_text_frame:
                text_boxes += 1
        
        print(f"   📈 Charts: {charts_on_slide}")
        print(f"   📋 Tables: {tables_on_slide}")
        print(f"   📝 Text Boxes: {text_boxes}")
        
        # Check for insights/data-driven content in Slide 2 (Executive Summary)
        if slide_idx == 2:
            print(f"\n   🎯 EXECUTIVE SUMMARY INSIGHTS:")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text
                    if '✓' in text or 'entities' in text.lower() or 'leads' in text.lower():
                        lines = [line.strip() for line in text.split('\n') if line.strip() and '✓' in line]
                        for line in lines[:3]:  # First 3 insights
                            print(f"      {line[:70]}...")
    
    # Summary
    print(f"\n{'='*80}")
    print("📊 OVERALL STATISTICS")
    print(f"{'='*80}")
    print(f"Total Charts: {total_charts}")
    print(f"Total Tables: {total_tables}")
    print(f"Unique Chart Types: {len(set(chart_types))}")
    
    if chart_types:
        print(f"\n📈 Chart Type Distribution:")
        from collections import Counter
        chart_counter = Counter(chart_types)
        for chart_type, count in chart_counter.items():
            print(f"   {chart_type}: {count}x")
    
    # Check for improvements
    print(f"\n{'='*80}")
    print("✅ IMPROVEMENTS VERIFIED:")
    print(f"{'='*80}")
    
    if total_charts >= 3:
        print(f"✅ Multiple diverse charts ({total_charts} total)")
    else:
        print(f"⚠️  Limited charts ({total_charts} total)")
    
    if len(set(chart_types)) >= 2:
        print(f"✅ Chart variety ({len(set(chart_types))} different types)")
    else:
        print(f"⚠️  Limited chart variety")
    
    if total_tables >= 1:
        print(f"✅ Data tables included ({total_tables} total)")
    
    print(f"\n{'='*80}")


if __name__ == "__main__":
    verify_improved_presentation()
