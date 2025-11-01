"""
Template-aware chart creation functions
Enhanced wrappers around ppt_writer.py functions with template support
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from src.templates.template_manager import TemplateManager, hex_to_rgb
from typing import Optional, Dict, Any, List


def create_templated_chart_slide(
    prs: Presentation,
    df: pd.DataFrame,
    chart_type: str,
    config: Dict[str, Any],
    title: str,
    template: Optional[Dict[str, Any]] = None
):
    """
    Create a chart slide with template styling
    
    Args:
        prs: Presentation object
        df: DataFrame with data
        chart_type: 'pie', 'bar', 'line', 'column', or 'scatter'
        config: Chart configuration dict
        title: Slide title
        template: Template configuration (optional)
    
    Returns:
        Created slide object
    """
    
    # Create blank slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Get template styling or use defaults
    if template:
        title_font = template['fonts']['heading']
        chart_colors = template['colors']['chart_colors']
        layout = template['layout']['content_slide']
        title_pos = layout['title_position']
        chart_pos = layout['chart_position']
    else:
        title_font = {'name': 'Calibri', 'size': 24, 'bold': True, 'color': '#333333'}
        chart_colors = ['#1F4788', '#4A90E2', '#7CB9E8', '#B8D4E8', '#2E5C8A']
        title_pos = {'left': 0.5, 'top': 0.2, 'width': 9, 'height': 0.5}
        chart_pos = {'left': 1, 'top': 1.2, 'width': 8, 'height': 5}
    
    # Add title
    title_box = slide.shapes.add_textbox(
        Inches(title_pos['left']),
        Inches(title_pos['top']),
        Inches(title_pos['width']),
        Inches(title_pos['height'])
    )
    title_box.text = title
    
    # Style title
    for paragraph in title_box.text_frame.paragraphs:
        paragraph.font.size = Pt(title_font['size'])
        paragraph.font.bold = title_font.get('bold', True)
        if 'color' in title_font:
            r, g, b = hex_to_rgb(title_font['color'])
            paragraph.font.color.rgb = RGBColor(r, g, b)
    
    # Create chart based on type
    if chart_type == 'pie':
        _create_pie_chart(slide, df, config, chart_pos, chart_colors)
    elif chart_type == 'bar':
        _create_bar_chart(slide, df, config, chart_pos, chart_colors)
    elif chart_type == 'line':
        _create_line_chart(slide, df, config, chart_pos, chart_colors)
    elif chart_type == 'column':
        _create_column_chart(slide, df, config, chart_pos, chart_colors)
    elif chart_type == 'scatter':
        _create_scatter_chart(slide, df, config, chart_pos, chart_colors)
    
    return slide


def _create_pie_chart(slide, df, config, chart_pos, colors):
    """Create pie chart with template colors"""
    category_col = config['category_col']
    value_col = config['value_col']
    
    # Prepare data
    chart_df = df[[category_col, value_col]].copy()
    chart_df = chart_df.dropna()
    chart_df = chart_df.sort_values(by=value_col, ascending=False)
    
    if len(chart_df) > 10:
        chart_df = chart_df.head(10)
    
    # Create chart data
    chart_data = CategoryChartData()
    chart_data.categories = chart_df[category_col].tolist()
    chart_data.add_series('Values', chart_df[value_col].tolist())
    
    # Add chart
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(chart_pos['left']),
        Inches(chart_pos['top']),
        Inches(chart_pos['width']),
        Inches(chart_pos['height']),
        chart_data
    ).chart
    
    # Apply colors
    _apply_chart_colors(chart, colors, is_pie=True)
    
    # Styling
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.font.size = Pt(10)
    
    chart.plots[0].has_data_labels = True
    data_labels = chart.plots[0].data_labels
    data_labels.number_format = '0%'
    data_labels.font.size = Pt(10)


def _create_bar_chart(slide, df, config, chart_pos, colors):
    """Create horizontal bar chart"""
    x_col = config['x_col']
    y_col = config['y_col']
    
    chart_df = df[[x_col, y_col]].copy().dropna()
    
    chart_data = CategoryChartData()
    chart_data.categories = chart_df[x_col].tolist()
    chart_data.add_series(y_col, chart_df[y_col].tolist())
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(chart_pos['left']),
        Inches(chart_pos['top']),
        Inches(chart_pos['width']),
        Inches(chart_pos['height']),
        chart_data
    ).chart
    
    _apply_chart_colors(chart, colors)
    
    chart.has_legend = False
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = '#,##0'


def _create_line_chart(slide, df, config, chart_pos, colors):
    """Create line chart for time series"""
    x_col = config['x_col']
    y_cols = config['y_cols']
    
    # Clean dataframe - remove NaN rows
    chart_df = df[[x_col] + y_cols].dropna()
    
    # Sample if too many points
    if len(chart_df) > 50:
        chart_df = chart_df.iloc[::max(1, len(chart_df) // 50)]
    
    chart_data = CategoryChartData()
    chart_data.categories = chart_df[x_col].astype(str).tolist()
    
    for i, y_col in enumerate(y_cols[:3]):  # Max 3 series
        if y_col in chart_df.columns:
            chart_data.add_series(y_col, chart_df[y_col].tolist())
    
    # Handle empty data
    if len(chart_df) == 0:
        return
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(chart_pos['left']),
        Inches(chart_pos['top']),
        Inches(chart_pos['width']),
        Inches(chart_pos['height']),
        chart_data
    ).chart
    
    _apply_chart_colors(chart, colors)
    
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(10)


def _create_column_chart(slide, df, config, chart_pos, colors):
    """Create vertical column chart"""
    x_col = config['x_col']
    y_cols = config['y_cols']
    
    # Clean dataframe - remove NaN rows
    chart_df = df[[x_col] + y_cols].dropna()
    
    if len(chart_df) == 0:
        return
    
    chart_data = CategoryChartData()
    chart_data.categories = chart_df[x_col].astype(str).tolist()
    
    for y_col in y_cols[:5]:  # Max 5 series
        if y_col in chart_df.columns:
            chart_data.add_series(y_col, chart_df[y_col].tolist())
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(chart_pos['left']),
        Inches(chart_pos['top']),
        Inches(chart_pos['width']),
        Inches(chart_pos['height']),
        chart_data
    ).chart
    
    _apply_chart_colors(chart, colors)
    
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(10)


def _create_scatter_chart(slide, df, config, chart_pos, colors):
    """Create scatter plot"""
    x_col = config['x_col']
    y_col = config['y_col']
    
    chart_df = df[[x_col, y_col]].copy().dropna()
    
    chart_data = XyChartData()
    series = chart_data.add_series(f'{y_col} vs {x_col}')
    
    for _, row in chart_df.iterrows():
        series.add_data_point(float(row[x_col]), float(row[y_col]))
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        Inches(chart_pos['left']),
        Inches(chart_pos['top']),
        Inches(chart_pos['width']),
        Inches(chart_pos['height']),
        chart_data
    ).chart
    
    _apply_chart_colors(chart, colors)


def _apply_chart_colors(chart, colors: List[str], is_pie: bool = False):
    """Apply template colors to chart series"""
    try:
        if is_pie:
            # For pie charts, color each point
            plot = chart.plots[0]
            for i, point in enumerate(plot.series[0].points):
                color_idx = i % len(colors)
                r, g, b = hex_to_rgb(colors[color_idx])
                fill = point.format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(r, g, b)
        else:
            # For other charts, color each series
            for i, series in enumerate(chart.series):
                color_idx = i % len(colors)
                r, g, b = hex_to_rgb(colors[color_idx])
                fill = series.format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(r, g, b)
    except Exception as e:
        print(f"Note: Could not apply custom colors: {e}")


def create_templated_data_table(
    slide,
    df: pd.DataFrame,
    template: Optional[Dict[str, Any]] = None,
    max_rows: int = 10
):
    """
    Create a styled data table on a slide
    
    Args:
        slide: Slide to add table to
        df: DataFrame to display
        template: Template configuration
        max_rows: Maximum rows to display
    """
    
    # Get template styling
    if template:
        header_color = template['styling']['table_header_color']
        alt_row_color = template['styling']['table_alt_row_color']
        table_pos = template['layout']['content_slide']['table_position']
        header_font = template['fonts']['table_header']
        data_font = template['fonts']['table_data']
    else:
        header_color = '#1F4788'
        alt_row_color = '#F2F2F2'
        table_pos = {'left': 6.2, 'top': 1.2, 'width': 3.5, 'height': 5}
        header_font = {'name': 'Calibri', 'size': 11, 'bold': True, 'color': '#FFFFFF'}
        data_font = {'name': 'Calibri', 'size': 10, 'bold': False, 'color': '#333333'}
    
    # Limit rows
    display_df = df.head(max_rows)
    
    rows = len(display_df) + 1
    cols = min(len(display_df.columns), 3)  # Max 3 columns in side table
    
    # Create table
    table = slide.shapes.add_table(
        rows, cols,
        Inches(table_pos['left']),
        Inches(table_pos['top']),
        Inches(table_pos['width']),
        Inches(table_pos['height'])
    ).table
    
    # Header row
    for col_idx in range(cols):
        cell = table.cell(0, col_idx)
        cell.text = str(display_df.columns[col_idx])[:15]
        
        # Header styling
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(header_font['size'])
            paragraph.font.bold = header_font.get('bold', True)
            if 'color' in header_font:
                r, g, b = hex_to_rgb(header_font['color'])
                paragraph.font.color.rgb = RGBColor(r, g, b)
        
        # Header background
        r, g, b = hex_to_rgb(header_color)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(r, g, b)
    
    # Data rows
    for row_idx, (_, row) in enumerate(display_df.iterrows(), start=1):
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            value = row.iloc[col_idx]
            
            # Format value
            if isinstance(value, float):
                if abs(value) >= 1_000_000_000:
                    cell.text = f"{value/1_000_000_000:.1f}B"
                elif abs(value) >= 1_000_000:
                    cell.text = f"{value/1_000_000:.1f}M"
                elif abs(value) >= 1000:
                    cell.text = f"{value/1000:.0f}K"
                else:
                    cell.text = f"{value:.1f}"
            else:
                cell.text = str(value)[:12]
            
            # Data font styling
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(data_font['size'])
                paragraph.font.bold = data_font.get('bold', False)
            
            # Alternate row colors
            if row_idx % 2 == 0:
                r, g, b = hex_to_rgb(alt_row_color)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(r, g, b)
    
    return table
