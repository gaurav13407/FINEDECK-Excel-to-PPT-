"""
Simple Finance-Optimized Chart Builder
NO AI, NO SmartChartAnalyzer, NO AdvancedChartBuilder
ONLY finance-specific chart detection and creation
"""

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pandas as pd
import numpy as np


class SimpleFinanceChartBuilder:
    """
    Simplified finance chart builder with NO dependencies on AI systems
    Uses only: PIE, COLUMN, LINE charts based on column names
    """
    
    def __init__(self, template_colors=None):
        """Initialize with optional template colors"""
        self.template_colors = template_colors or {
            'primary': '#1E3A8A',
            'secondary': '#3B82F6', 
            'accent': '#60A5FA',
            'success': '#10B981',
            'warning': '#F59E0B',
            'danger': '#EF4444'
        }
        
        # Finance-specific chart colors
        self.chart_colors = [
            RGBColor(30, 58, 138),   # Deep Blue
            RGBColor(59, 130, 246),  # Blue
            RGBColor(96, 165, 250),  # Light Blue
            RGBColor(16, 185, 129),  # Green
            RGBColor(245, 158, 11),  # Orange
            RGBColor(239, 68, 68),   # Red
            RGBColor(139, 92, 246),  # Purple
            RGBColor(236, 72, 153),  # Pink
        ]
    
    def detect_chart_type(self, df):
        """
        Detect best chart type based on column names
        Finance-optimized priority order:
        1. TIME-SERIES (date/quarter/month/year) → LINE
        2. ALLOCATION (allocation/portfolio/sector) → PIE
        3. PERFORMANCE (top/rank/performance/revenue) → COLUMN
        4. DEFAULT → COLUMN (finance standard)
        """
        print("\n🔍 DETECTING CHART TYPE...")
        print(f"   Columns: {list(df.columns)}")
        
        col_names_lower = ' '.join([str(col).lower() for col in df.columns])
        print(f"   Column names (lowercase): '{col_names_lower}'")
        
        # PRIORITY 1: Time-series data → LINE chart
        time_keywords = ['date', 'quarter', 'month', 'year', 'q1', 'q2', 'q3', 'q4', 'ytd', 'period']
        if any(kw in col_names_lower for kw in time_keywords):
            matched = [kw for kw in time_keywords if kw in col_names_lower]
            print(f"   ⏰ Time-series keywords matched: {matched}")
            print(f"   → Returning: LINE chart")
            return 'line'
        
        # PRIORITY 2: Allocation/distribution data → PIE chart
        allocation_keywords = ['allocation', 'portfolio', 'sector', 'distribution', 'breakdown', 'composition']
        if any(kw in col_names_lower for kw in allocation_keywords):
            matched = [kw for kw in allocation_keywords if kw in col_names_lower]
            print(f"   🥧 Allocation keywords matched: {matched}")
            print(f"   → Returning: PIE chart")
            return 'pie'
        
        # PRIORITY 3: Performance/ranking data → COLUMN chart
        performance_keywords = ['top', 'rank', 'performance', 'revenue', 'sales', 'profit', 'growth', 'comparison']
        if any(kw in col_names_lower for kw in performance_keywords):
            matched = [kw for kw in performance_keywords if kw in col_names_lower]
            print(f"   📊 Performance keywords matched: {matched}")
            print(f"   → Returning: COLUMN chart")
            return 'column'
        
        # DEFAULT: COLUMN chart (finance standard for comparisons)
        print(f"   ⚠️  No keywords matched")
        print(f"   → Returning: COLUMN chart (default)")
        return 'column'
    
    def create_chart(self, slide, df, left, top, width, height, title=None):
        """
        Create a finance-appropriate chart
        
        Args:
            slide: PowerPoint slide object
            df: DataFrame with data
            left, top, width, height: Position in EMUs (from Inches())
            title: Optional chart title
        
        Returns:
            Chart object or None
        """
        print("\n" + "="*80)
        print("🎨 SIMPLE FINANCE CHART BUILDER - CREATE_CHART()")
        print("="*80)
        print(f"📊 DataFrame shape: {df.shape if df is not None else 'None'}")
        print(f"📝 Title: {title}")
        print(f"📍 Position: left={left} EMUs ({left/914400:.2f}\"), top={top} EMUs ({top/914400:.2f}\")")
        print(f"📏 Size: width={width} EMUs ({width/914400:.2f}\"), height={height} EMUs ({height/914400:.2f}\")")
        
        if df is None or df.empty:
            print("   ❌ Cannot create chart: DataFrame is None or empty")
            print("="*80 + "\n")
            return None
        
        print(f"📋 Columns: {list(df.columns)}")
        
        # Clean DataFrame: Remove NaN/Inf values
        df = df.replace([np.inf, -np.inf], np.nan).fillna(0)
        
        # Get text and numeric columns
        text_cols = df.select_dtypes(include=['object']).columns
        numeric_cols = df.select_dtypes(include=['number']).columns
        
        print(f"📊 Text columns: {list(text_cols)}")
        print(f"🔢 Numeric columns: {list(numeric_cols)}")
        
        if len(numeric_cols) == 0:
            print("   ❌ Cannot create chart: No numeric columns")
            print("="*80 + "\n")
            return None
        
        # Detect chart type
        chart_type = self.detect_chart_type(df)
        
        print(f"🎯 DETECTED CHART TYPE: {chart_type.upper()}")
        print("="*80)
        
        # Create chart based on type
        if chart_type == 'pie':
            result = self._create_pie_chart(slide, df, left, top, width, height, title, text_cols, numeric_cols)
        elif chart_type == 'line':
            result = self._create_line_chart(slide, df, left, top, width, height, title, text_cols, numeric_cols)
        else:  # column or default
            result = self._create_column_chart(slide, df, left, top, width, height, title, text_cols, numeric_cols)
        
        if result:
            print(f"✅ Chart created successfully: {chart_type.upper()}")
        else:
            print(f"❌ Chart creation failed")
        print("="*80 + "\n")
        
        return result
    
    def _create_pie_chart(self, slide, df, left, top, width, height, title, text_cols, numeric_cols):
        """Create a PIE chart for allocation/distribution data"""
        print("\n   🥧 CREATING PIE CHART...")
        try:
            chart_data = CategoryChartData()
            
            # SMART: Find best label column (avoid dates, prefer sector/category names)
            label_col = None
            for col in text_cols:
                col_lower = str(col).lower()
                # Skip date-like columns
                if not any(kw in col_lower for kw in ['date', 'year', 'month', 'quarter']):
                    label_col = col
                    break
            
            # Fallback to first text column or index
            if not label_col:
                label_col = text_cols[0] if len(text_cols) > 0 else df.index.name
            
            # SMART: Find best value column (prefer meaningful financial metrics)
            value_col = None
            for col in numeric_cols[:10]:  # Check first 10 numeric columns
                col_lower = str(col).lower()
                # Prefer columns with value-like names
                if any(kw in col_lower for kw in ['value', 'amount', 'total', 'revenue', 'sales', 'allocation']):
                    value_col = col
                    break
            
            # Fallback to first numeric column
            if not value_col and len(numeric_cols) > 0:
                value_col = numeric_cols[0]
            
            if not value_col:
                print("      ❌ No numeric column found for values")
                return None
            
            print(f"      Label column: {label_col}")
            print(f"      Value column: {value_col}")
            
            # Get data (limit to top 8 for readability)
            data_df = df.head(8).copy()
            print(f"      Data rows: {len(data_df)}")
            
            # Add categories and values with NaN cleaning
            chart_data.categories = [str(cat) for cat in data_df[label_col]] if label_col else [str(i) for i in range(len(data_df))]
            
            # Clean values
            values = []
            for val in data_df[value_col]:
                try:
                    if pd.isna(val) or np.isnan(float(val)) or np.isinf(float(val)):
                        values.append(0.0)
                    else:
                        values.append(float(val))
                except (ValueError, TypeError):
                    values.append(0.0)
            
            chart_data.add_series('Values', values)
            
            print(f"      Categories: {len(chart_data.categories)}")
            print(f"      Values: {values[:3]}... (showing first 3)")
            
            # Create chart - USE EMU VALUES DIRECTLY (already from Inches())
            print(f"      📍 Position: left={left} EMUs ({left/914400:.2f}\"), top={top} EMUs ({top/914400:.2f}\")")
            print(f"      📏 Size: {width} EMUs ({width/914400:.2f}\") x {height} EMUs ({height/914400:.2f}\")")
            
            chart_placeholder = slide.shapes.add_chart(
                XL_CHART_TYPE.PIE,
                left, top, width, height,  # Already in EMUs - DO NOT wrap in Inches()
                chart_data
            )
            
            print(f"      ✅ PIE chart created with {len(chart_data.categories)} slices")
            
            chart = chart_placeholder.chart
            if title:
                chart.has_title = True
                chart.chart_title.text_frame.text = title
            
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.RIGHT
            
            # Apply colors
            for idx, point in enumerate(chart.plots[0].series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.chart_colors[idx % len(self.chart_colors)]
            
            print(f"   ✅ PIE chart created successfully with {len(values)} slices")
            return chart
            
        except Exception as e:
            print(f"   ❌ Error creating PIE chart: {e}")
            return None
    
    def _create_line_chart(self, slide, df, left, top, width, height, title, text_cols, numeric_cols):
        """Create a LINE chart for time-series data"""
        print("\n   📈 CREATING LINE CHART...")
        try:
            chart_data = CategoryChartData()
            
            # SMART: Look for Date column for X-axis (time-series should use dates!)
            date_col = None
            for col in text_cols:
                col_lower = str(col).lower()
                if any(kw in col_lower for kw in ['date', 'quarter', 'month', 'year', 'period']):
                    date_col = col
                    break
            
            # Use date column if found, otherwise first text column
            if date_col:
                categories = [str(cat) for cat in df[date_col]]
                print(f"      X-axis column: {date_col} (detected as time column)")
            elif len(text_cols) > 0:
                categories = [str(cat) for cat in df[text_cols[0]]]
                print(f"      X-axis column: {text_cols[0]} (first text column)")
            else:
                categories = [str(cat) for cat in df.index]
                print(f"      X-axis: Using index")
            
            chart_data.categories = categories
            
            print(f"      Categories: {len(categories)}")
            
            # Add up to 3 numeric series (skip date-like numeric columns)
            series_cols = []
            for col in numeric_cols[:10]:  # Check first 10 numeric columns
                col_lower = str(col).lower()
                # Skip columns that look like dates or IDs
                if not any(kw in col_lower for kw in ['date', 'id', 'year', 'month']):
                    series_cols.append(col)
                    if len(series_cols) >= 3:
                        break
            
            if not series_cols:  # Fallback to first 3 numeric columns
                series_cols = numeric_cols[:3]
            
            print(f"      Y-axis columns: {list(series_cols)}")
            
            for col in series_cols:
                # Clean values
                values = []
                for val in df[col]:
                    try:
                        if pd.isna(val) or np.isnan(float(val)) or np.isinf(float(val)):
                            values.append(0.0)
                        else:
                            values.append(float(val))
                    except (ValueError, TypeError):
                        values.append(0.0)
                
                chart_data.add_series(str(col), values)
            
            # Create chart - USE EMU VALUES DIRECTLY
            print(f"      📍 Position: left={left} EMUs ({left/914400:.2f}\"), top={top} EMUs ({top/914400:.2f}\")")
            print(f"      📏 Size: {width} EMUs ({width/914400:.2f}\") x {height} EMUs ({height/914400:.2f}\")")
            
            chart_placeholder = slide.shapes.add_chart(
                XL_CHART_TYPE.LINE_MARKERS,
                left, top, width, height,  # Already in EMUs
                chart_data
            )
            
            print(f"      ✅ LINE chart created with {len(series_cols)} series")
            
            chart = chart_placeholder.chart
            if title:
                chart.has_title = True
                chart.chart_title.text_frame.text = title
            
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            
            # Style series
            for idx, series in enumerate(chart.plots[0].series):
                series.format.line.color.rgb = self.chart_colors[idx % len(self.chart_colors)]
                series.format.line.width = Pt(2)
                # Set marker style using valid enum
                series.marker.style = XL_MARKER_STYLE.CIRCLE
            
            print(f"   ✅ LINE chart created with {len(series_cols)} series")
            return chart
            
        except Exception as e:
            print(f"   ❌ Error creating LINE chart: {e}")
            return None
    
    def _create_column_chart(self, slide, df, left, top, width, height, title, text_cols, numeric_cols):
        """Create a COLUMN chart for performance/comparison data"""
        print("\n   📊 CREATING COLUMN CHART...")
        try:
            chart_data = CategoryChartData()
            
            # SMART: Find best label column (avoid dates)
            label_col = None
            for col in text_cols:
                col_lower = str(col).lower()
                # Skip date-like columns
                if not any(kw in col_lower for kw in ['date', 'year', 'month', 'quarter']):
                    label_col = col
                    break
            
            # Use label column if found, otherwise first text column or index
            if label_col:
                categories = [str(cat) for cat in df[label_col].head(10)]
                print(f"      X-axis column: {label_col} (detected as category)")
            elif len(text_cols) > 0:
                categories = [str(cat) for cat in df[text_cols[0]].head(10)]
                print(f"      X-axis column: {text_cols[0]} (first text column)")
            else:
                categories = [str(cat) for cat in df.index[:10]]
                print(f"      X-axis: Using index")
            
            chart_data.categories = categories
            
            print(f"      Categories: {len(categories)}")
            
            # SMART: Find best value column (prefer meaningful metrics)
            value_col = None
            for col in numeric_cols[:10]:  # Check first 10 numeric columns
                col_lower = str(col).lower()
                # Prefer columns with value-like names, skip date/ID columns
                if not any(kw in col_lower for kw in ['date', 'id', 'year']):
                    if any(kw in col_lower for kw in ['value', 'amount', 'revenue', 'sales', 'profit', 'performance', 'total']):
                        value_col = col
                        break
            
            # Fallback to first non-date numeric column
            if not value_col:
                for col in numeric_cols[:10]:
                    col_lower = str(col).lower()
                    if not any(kw in col_lower for kw in ['date', 'year', 'id']):
                        value_col = col
                        break
            
            # Last resort: first numeric column
            if not value_col and len(numeric_cols) > 0:
                value_col = numeric_cols[0]
            
            if not value_col:
                print("      ❌ No numeric column found for values")
                return None
            
            print(f"      Y-axis column: {value_col}")
            
            values = []
            for val in df[value_col].head(10):
                try:
                    if pd.isna(val) or np.isnan(float(val)) or np.isinf(float(val)):
                        values.append(0.0)
                    else:
                        values.append(float(val))
                except (ValueError, TypeError):
                    values.append(0.0)
            
            chart_data.add_series(str(value_col), values)
            
            # Create chart - USE EMU VALUES DIRECTLY
            print(f"      📍 Position: left={left} EMUs ({left/914400:.2f}\"), top={top} EMUs ({top/914400:.2f}\")")
            print(f"      📏 Size: {width} EMUs ({width/914400:.2f}\") x {height} EMUs ({height/914400:.2f}\")")
            
            chart_placeholder = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                left, top, width, height,  # Already in EMUs
                chart_data
            )
            
            print(f"      ✅ COLUMN chart created with {len(categories)} bars")
            
            chart = chart_placeholder.chart
            if title:
                chart.has_title = True
                chart.chart_title.text_frame.text = title
            
            chart.has_legend = False  # Single series doesn't need legend
            
            # Apply colors
            series = chart.plots[0].series[0]
            for idx, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.chart_colors[idx % len(self.chart_colors)]
            
            print(f"   ✅ COLUMN chart created with {len(values)} bars")
            return chart
            
        except Exception as e:
            print(f"   ❌ Error creating COLUMN chart: {e}")
            return None
