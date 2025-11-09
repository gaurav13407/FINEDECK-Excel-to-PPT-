"""
Diagnostic Tool: Check what charts are being created in your PPTs
This will analyze the most recent PPT files and show detailed chart information
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

def analyze_ppt_charts(ppt_path):
    """Analyze all charts in a PowerPoint file"""
    print(f"\n{'='*80}")
    print(f"📊 ANALYZING: {os.path.basename(ppt_path)}")
    print(f"{'='*80}")
    
    try:
        prs = Presentation(ppt_path)
        total_charts = 0
        
        for slide_idx, slide in enumerate(prs.slides, 1):
            slide_title = "Untitled"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and len(text) < 100:
                        slide_title = text
                        break
            
            charts_on_slide = []
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = shape.chart
                    total_charts += 1
                    
                    # Get chart type name
                    chart_type = chart.chart_type
                    chart_type_name = str(chart_type).split('.')[-1].replace('(', '').replace(')', '').split()[0]
                    chart_type_value = int(str(chart_type).split('(')[-1].split(')')[0])
                    
                    # Determine what type it is
                    chart_category = "UNKNOWN"
                    if "PIE" in chart_type_name or chart_type_value == 5:
                        chart_category = "🥧 PIE CHART"
                    elif "BAR" in chart_type_name or chart_type_value == 57:
                        chart_category = "📊 BAR CHART (Horizontal)"
                    elif "COLUMN" in chart_type_name or chart_type_value == 51:
                        chart_category = "📊 COLUMN CHART (Vertical)"
                    elif "LINE" in chart_type_name or chart_type_value == 65:
                        chart_category = "📈 LINE CHART"
                    elif "DOUGHNUT" in chart_type_name:
                        chart_category = "🍩 DOUGHNUT CHART"
                    
                    # Get chart data
                    series_count = len(chart.plots[0].series) if chart.plots and chart.plots[0].series else 0
                    categories_count = 0
                    if chart.plots and chart.plots[0].categories:
                        categories_count = len(chart.plots[0].categories)
                    
                    # Get colors
                    colors = []
                    try:
                        for series in chart.plots[0].series:
                            if hasattr(series, 'format') and hasattr(series.format, 'fill'):
                                fill = series.format.fill
                                if hasattr(fill, 'fore_color') and hasattr(fill.fore_color, 'rgb'):
                                    rgb = fill.fore_color.rgb
                                    colors.append(f"RGB{rgb}")
                    except:
                        pass
                    
                    chart_info = {
                        'category': chart_category,
                        'type_name': chart_type_name,
                        'type_value': chart_type_value,
                        'series': series_count,
                        'categories': categories_count,
                        'colors': colors
                    }
                    charts_on_slide.append(chart_info)
            
            if charts_on_slide:
                print(f"\n📄 Slide {slide_idx}: {slide_title[:60]}")
                for i, chart_info in enumerate(charts_on_slide, 1):
                    print(f"   Chart #{i}:")
                    print(f"      Type: {chart_info['category']}")
                    print(f"      Technical: {chart_info['type_name']} ({chart_info['type_value']})")
                    print(f"      Data: {chart_info['series']} series, {chart_info['categories']} categories")
                    if chart_info['colors']:
                        print(f"      Colors: {', '.join(chart_info['colors'][:3])}")
        
        print(f"\n{'='*80}")
        print(f"📊 SUMMARY: Found {total_charts} chart(s) in this presentation")
        print(f"{'='*80}\n")
        
        return total_charts
        
    except Exception as e:
        print(f"❌ Error analyzing {ppt_path}: {e}")
        import traceback
        traceback.print_exc()
        return 0

def find_recent_ppts(directory, limit=5):
    """Find most recent PPT files"""
    ppt_files = []
    
    for ext in ['*.pptx', '*.ppt']:
        ppt_files.extend(Path(directory).rglob(ext))
    
    # Sort by modification time (most recent first)
    ppt_files.sort(key=lambda x: x.stat().st_mtime, reverse=True)
    
    return ppt_files[:limit]

def main():
    print("\n" + "="*80)
    print("🔍 CHART TYPE DIAGNOSTIC TOOL")
    print("="*80)
    print("\nThis tool will analyze your PPT files and show:")
    print("  - What types of charts are being created")
    print("  - Whether they're PIE, BAR (horizontal), COLUMN (vertical), or LINE")
    print("  - Chart colors")
    print("  - Data structure (series, categories)")
    
    # Check common output directories
    search_dirs = [
        r"C:\Users\gaura\OneDrive\Desktop\Big Projects\FinDeck(Excel to PPT Project)\test_output",
        r"C:\Users\gaura\OneDrive\Desktop\Big Projects\FinDeck(Excel to PPT Project)\examples\demo_PPT",
        r"C:\Users\gaura\OneDrive\Desktop\Big Projects\FinDeck(Excel to PPT Project)",
    ]
    
    all_ppts = []
    for dir_path in search_dirs:
        if os.path.exists(dir_path):
            all_ppts.extend(find_recent_ppts(dir_path, limit=10))
    
    # Remove duplicates and sort by modification time
    all_ppts = list(set(all_ppts))
    all_ppts.sort(key=lambda x: x.stat().st_mtime, reverse=True)
    
    if not all_ppts:
        print("\n❌ No PowerPoint files found!")
        print("\nPlease specify the path to your PPT file:")
        ppt_path = input("PPT Path: ").strip('"').strip("'")
        if os.path.exists(ppt_path):
            analyze_ppt_charts(ppt_path)
        else:
            print(f"❌ File not found: {ppt_path}")
        return
    
    print(f"\n📁 Found {len(all_ppts)} recent PPT file(s):")
    for i, ppt in enumerate(all_ppts[:10], 1):
        mod_time = datetime.fromtimestamp(ppt.stat().st_mtime)
        print(f"   {i}. {ppt.name} (Modified: {mod_time.strftime('%Y-%m-%d %H:%M:%S')})")
    
    print("\n" + "="*80)
    print("ANALYZING THE 5 MOST RECENT FILES...")
    print("="*80)
    
    total_analyzed = 0
    for ppt in all_ppts[:5]:
        chart_count = analyze_ppt_charts(str(ppt))
        total_analyzed += 1
    
    print("\n" + "="*80)
    print("🎯 WHAT CHARTS DO YOU WANT?")
    print("="*80)
    print("\nCurrently, the auto-detection creates:")
    print("  📊 BAR charts (horizontal) - for small datasets (≤10 rows)")
    print("  📊 COLUMN charts (vertical) - for large datasets (>10 rows)")
    print("  🥧 PIE charts - for data with '%' or 'percent' columns")
    print("  📈 LINE charts - for data with date/time columns")
    
    print("\n💡 To change this behavior, we can:")
    print("  1. Adjust the auto-detection thresholds")
    print("  2. Prefer specific chart types (e.g., always use PIE charts)")
    print("  3. Add user selection in the UI")
    print("  4. Create custom rules based on data patterns")
    
    print("\n📝 Please tell me:")
    print("  - What chart type do you WANT to see?")
    print("  - For what kind of data?")
    print("  - Should it be automatic or user-selected?")
    
    print("\n" + "="*80)

if __name__ == "__main__":
    main()
