"""
Direct test of chart creation to verify it works
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'src'))

from pptx import Presentation
from pptx.util import Inches
import pandas as pd
from converter.enhanced_charts import EnhancedChartBuilder

print("=" * 80)
print("DIRECT CHART CREATION TEST")
print("=" * 80)

# Create a simple presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Add a blank slide
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Create test data
test_data = pd.DataFrame({
    'Category': ['A', 'B', 'C', 'D', 'E'],
    'Value': [100, 200, 150, 300, 250]
})

print("\n📊 Test Data:")
print(test_data)

# Initialize EnhancedChartBuilder
template_colors = {
    'navy': (31, 73, 125),
    'light_blue': (68, 114, 196),
    'teal': (91, 155, 213),
    'gray': (165, 165, 165),
    'dark_gray': (89, 89, 89),
    'accent_orange': (237, 125, 49),
    'accent_green': (112, 173, 71)
}

builder = EnhancedChartBuilder(template_colors)

print("\n🎨 Creating chart with auto_create_chart...")
print(f"   Position: left=Inches(1), top=Inches(1), width=Inches(8), height=Inches(5)")

# Create chart
chart = builder.auto_create_chart(
    slide=slide,
    df=test_data,
    left=Inches(1),
    top=Inches(1),
    width=Inches(8),
    height=Inches(5),
    title="Test Chart"
)

if chart:
    print("   ✅ Chart created successfully!")
    print(f"   Chart type: {chart.chart_type}")
else:
    print("   ❌ Chart creation returned None!")

# Check if chart was added to slide
chart_count = sum(1 for shape in slide.shapes if hasattr(shape, 'chart'))
print(f"\n📊 Charts on slide: {chart_count}")

if chart_count > 0:
    for i, shape in enumerate(slide.shapes):
        if hasattr(shape, 'chart'):
            print(f"   Chart {i+1}:")
            print(f"     Position: left={shape.left}, top={shape.top}")
            print(f"     Size: width={shape.width}, height={shape.height}")
            print(f"     Left in inches: {shape.left / 914400}")
            print(f"     Top in inches: {shape.top / 914400}")

# Save presentation
output_path = "test_chart_output.pptx"
prs.save(output_path)
print(f"\n💾 Saved to: {output_path}")

print("\n" + "=" * 80)
print("VERIFICATION")
print("=" * 80)

# Reload and verify
prs2 = Presentation(output_path)
slide2 = prs2.slides[0]
chart_count2 = sum(1 for shape in slide2.shapes if hasattr(shape, 'chart'))

print(f"Charts in saved file: {chart_count2}")

if chart_count2 > 0:
    print("✅ SUCCESS! Chart was saved to the file!")
    print("\nOpen test_chart_output.pptx to see the chart")
else:
    print("❌ FAILED! Chart not found in saved file!")

print("\nExpected position in EMUs:")
print(f"  left=Inches(1) = {Inches(1)} EMUs")
print(f"  top=Inches(1) = {Inches(1)} EMUs")
print(f"  width=Inches(8) = {Inches(8)} EMUs")
print(f"  height=Inches(5) = {Inches(5)} EMUs")

if chart_count2 > 0:
    for shape in slide2.shapes:
        if hasattr(shape, 'chart'):
            print(f"\nActual position in saved file:")
            print(f"  left = {shape.left} EMUs ({shape.left / 914400:.2f} inches)")
            print(f"  top = {shape.top} EMUs ({shape.top / 914400:.2f} inches)")
            print(f"  width = {shape.width} EMUs ({shape.width / 914400:.2f} inches)")
            print(f"  height = {shape.height} EMUs ({shape.height / 914400:.2f} inches)")
            
            if shape.left == Inches(1):
                print("  ✅ Position is CORRECT!")
            else:
                print("  ❌ Position is WRONG!")
