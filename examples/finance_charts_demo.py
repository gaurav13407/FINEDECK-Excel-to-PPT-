"""
Finance Charts Demo
Generates a single PPTX containing multiple finance-appropriate charts using
AdvancedChartBuilder with context='finance'.

Run:
    python examples/finance_charts_demo.py

Outputs:
    examples/professional_demo/finance_charts_demo.pptx
"""

import os
from pathlib import Path
import sys
import pandas as pd
import numpy as np
from pptx import Presentation
from pptx.util import Inches

# Ensure project root is importable when running this script directly
project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from src.converter.advanced_chart_templates import AdvancedChartBuilder

OUT_DIR = Path("examples/professional_demo")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT_DIR / "finance_charts_demo.pptx"

builder = AdvancedChartBuilder(ai_service=None, smart_analyzer=None)

prs = Presentation()

# -----------------------
# 1) Trend (Line) chart
# -----------------------
trend_df = pd.DataFrame({
    'Period': pd.date_range(end=pd.Timestamp.today(), periods=12, freq='ME'),
    'Revenue': np.linspace(1000, 1500, 12) + np.random.randn(12) * 50,
    'Profit': np.linspace(200, 320, 12) + np.random.randn(12) * 10,
})
trend_df['Period'] = trend_df['Period'].dt.strftime('%Y-%m')
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart = builder.create_chart(
    slide, trend_df, builder.select_chart_type(trend_df, context='finance'),
    Inches(0.5), Inches(1.0), Inches(9), Inches(5.0),
    title="Revenue & Profit Trend (Line Chart)"
)

# -----------------------
# 2) Comparison (Column) chart
# -----------------------
comp_df = pd.DataFrame({
    'Product': [f"Product {i+1}" for i in range(8)],
    'Revenue': [1200, 980, 1500, 760, 430, 2100, 1750, 940]
})
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart = builder.create_chart(
    slide, comp_df, builder.select_chart_type(comp_df, context='finance'),
    Inches(0.5), Inches(1.0), Inches(9), Inches(5.0),
    title="Product Revenue Comparison (Column Chart)"
)

# -----------------------
# 3) Distribution (Doughnut)
# -----------------------
dist_df = pd.DataFrame({
    'Sector': ['Technology', 'Healthcare', 'Financials', 'Energy', 'Consumer', 'Utilities'],
    'MarketShare': [35, 18, 20, 8, 12, 7]
})
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart = builder.create_chart(
    slide, dist_df, builder.select_chart_type(dist_df, context='finance'),
    Inches(0.5), Inches(1.0), Inches(9), Inches(5.0),
    title="Market Share Distribution (Doughnut Chart)"
)

# -----------------------
# 4) Ranking (Bar) chart
# -----------------------
rank_df = pd.DataFrame({
    'Region': ['North America', 'Europe', 'Asia Pacific', 'Latin America', 'Middle East', 'Africa'],
    'Sales': [2500, 1800, 3200, 950, 600, 450]
})
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart = builder.create_chart(
    slide, rank_df, builder.select_chart_type(rank_df, context='finance'),
    Inches(0.5), Inches(1.0), Inches(9), Inches(5.0),
    title="Regional Sales Ranking (Bar Chart)"
)

# -----------------------
# 5) Portfolio Performance (Multi-line)
# -----------------------
portfolio_df = pd.DataFrame({
    'Quarter': ['Q1 2024', 'Q2 2024', 'Q3 2024', 'Q4 2024', 'Q1 2025'],
    'Stock_A': [100, 105, 110, 108, 115],
    'Stock_B': [100, 98, 103, 107, 110],
    'Stock_C': [100, 102, 99, 105, 108]
})
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart = builder.create_chart(
    slide, portfolio_df, builder.select_chart_type(portfolio_df, context='finance'),
    Inches(0.5), Inches(1.0), Inches(9), Inches(5.0),
    title="Portfolio Performance (Multi-Line Chart)"
)

# Save
prs.save(str(OUT_FILE))
print(f"Saved demo PPT to: {OUT_FILE}")
