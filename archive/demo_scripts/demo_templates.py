"""
Professional Template System Demo
Shows all 10 templates and custom template capabilities
"""

import sys
import os
sys.path.insert(0, 'src')

from converter.excel_reader import excel_reader
from pptx import Presentation
from pptx.util import Inches, Pt
from templates.template_manager import TemplateManager, get_default_template
from templates.template_charts import create_templated_chart_slide, create_templated_data_table
from converter.chart_detector import detect_chart_type, should_create_chart


def create_presentation_with_template(excel_file, output_file, template_name="corporate_blue"):
    """Create a professional presentation using a template"""
    
    print(f"\n{'='*70}")
    print(f"🎨 Creating Presentation with: {template_name.upper().replace('_', ' ')}")
    print(f"{'='*70}\n")
    
    # Load template
    manager = TemplateManager()
    template = manager.load_template(template_name)
    
    if not template:
        print(f"❌ Template '{template_name}' not found!")
        return None
    
    print(f"✅ Template loaded: {template['name']}")
    print(f"   Category: {template.get('category', 'N/A')}")
    print(f"   Description: {template['description'][:60]}...")
    
    # Create presentation
    prs = Presentation()
    
    # Create title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Financial Analysis"
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = f"Using {template['name']} Template"
    
    # Style title slide
    title_font = template['fonts']['title']
    subtitle_font = template['fonts']['subtitle']
    
    from pptx.dml.color import RGBColor
    from templates.template_manager import hex_to_rgb
    
    for paragraph in title_slide.shapes.title.text_frame.paragraphs:
        paragraph.font.size = Pt(title_font['size'])
        paragraph.font.bold = title_font.get('bold', True)
        if 'color' in title_font:
            r, g, b = hex_to_rgb(title_font['color'])
            paragraph.font.color.rgb = RGBColor(r, g, b)
    
    if len(title_slide.placeholders) > 1:
        for paragraph in title_slide.placeholders[1].text_frame.paragraphs:
            paragraph.font.size = Pt(subtitle_font['size'])
            if 'color' in subtitle_font:
                r, g, b = hex_to_rgb(subtitle_font['color'])
                paragraph.font.color.rgb = RGBColor(r, g, b)
    
    # Read data
    df = excel_reader(excel_file)
    
    if df is None or df.empty:
        print("   ⚠️  No data found")
        return None
    
    # Sample if too large
    if df.shape[0] > 30:
        df = df.tail(20)
    
    # Check if chartable
    if not should_create_chart(df):
        print("   ⚠️  Data not suitable for charts")
        return None
    
    # Detect chart type
    chart_type, config = detect_chart_type(df)
    
    if chart_type is None:
        print("   ⚠️  No chart type detected")
        return None
    
    print(f"   📊 Chart Type: {chart_type.upper()}")
    
    # Create chart slide with template
    slide = create_templated_chart_slide(
        prs, df, chart_type, config,
        "Financial Data Analysis",
        template=template
    )
    
    if slide:
        # Add data table to the same slide
        create_templated_data_table(slide, df, template=template, max_rows=8)
        print(f"   ✅ Chart + Data Table created!")
    
    # Save
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    prs.save(output_file)
    
    print(f"   💾 Saved: {output_file}")
    print(f"{'='*70}\n")
    
    return output_file


def demo_all_templates():
    """Create presentations with all 10 templates"""
    
    print("\n" + "="*70)
    print("🎨 PROFESSIONAL TEMPLATE SYSTEM DEMO")
    print("   Generating 10 presentations with different templates")
    print("="*70)
    
    # Initialize manager
    manager = TemplateManager()
    
    # List all templates
    print("\n📋 Available Templates:")
    print("-"*70)
    templates = manager.list_templates(include_custom=False)
    
    for i, template_info in enumerate(templates, 1):
        print(f"{i:2}. {template_info['name']:25} [{template_info['category']:12}]")
        print(f"    {template_info['description'][:60]}...")
    
    print(f"\n✅ Total: {len(templates)} built-in templates")
    print("="*70 + "\n")
    
    # Sample data file
    excel_file = "example/Company_Data/AAPL_Financial_Data.xlsx"
    
    if not os.path.exists(excel_file):
        print(f"❌ Sample file not found: {excel_file}")
        print("   Using Portfolio Allocation Data instead...")
        excel_file = "examples/Portfolio Allocation Data.xlsx"
    
    # Create presentations with each template
    created_files = []
    
    for template_info in templates:
        template_id = template_info['id']
        output_file = f"examples/demo_PPT/template_{template_id}.pptx"
        
        result = create_presentation_with_template(
            excel_file, 
            output_file, 
            template_id
        )
        
        if result:
            created_files.append((template_info['name'], result))
    
    # Summary
    print("\n" + "="*70)
    print("✅ DEMO COMPLETE!")
    print("="*70)
    print(f"\n📊 Created {len(created_files)} presentations:\n")
    
    for name, file_path in created_files:
        file_size = os.path.getsize(file_path) / 1024  # KB
        print(f"   • {name:25} → {os.path.basename(file_path):30} ({file_size:.1f} KB)")
    
    print("\n" + "="*70)
    print(f"📂 All files saved in: examples/demo_PPT/")
    print("="*70 + "\n")


def demo_custom_template():
    """Demo custom template creation"""
    
    print("\n" + "="*70)
    print("🎨 CUSTOM TEMPLATE DEMO")
    print("="*70 + "\n")
    
    manager = TemplateManager()
    
    # Create a custom template by duplicating
    print("1️⃣  Duplicating 'corporate_blue' to create custom template...")
    success = manager.duplicate_template("corporate_blue", "my_custom_theme")
    
    if success:
        print("   ✅ Custom template created: my_custom_theme")
        
        # Load and modify
        print("\n2️⃣  Loading custom template...")
        custom_template = manager.load_template("my_custom_theme")
        
        if custom_template:
            print("   ✅ Template loaded")
            print(f"   Name: {custom_template['name']}")
            print(f"   Primary Color: {custom_template['colors']['primary']}")
            
            # Modify colors
            print("\n3️⃣  Modifying template colors...")
            custom_template['name'] = "My Custom Business Theme"
            custom_template['colors']['primary'] = "#8B0000"  # Dark red
            custom_template['colors']['secondary'] = "#DC143C"  # Crimson
            custom_template['colors']['chart_colors'] = [
                "#8B0000", "#DC143C", "#FF6347", "#FFA07A", "#FFB6C1"
            ]
            
            # Save modified template
            manager.save_custom_template(custom_template, "my_custom_theme", overwrite=True)
            print("   ✅ Template colors updated")
            
            # Create presentation with custom template
            print("\n4️⃣  Creating presentation with custom template...")
            excel_file = "examples/Portfolio Allocation Data.xlsx"
            output_file = "examples/demo_PPT/template_my_custom_theme.pptx"
            
            result = create_presentation_with_template(
                excel_file,
                output_file,
                "my_custom_theme"
            )
            
            if result:
                print(f"   ✅ Custom template presentation created!")
                print(f"   💾 {output_file}")
    
    print("\n" + "="*70)
    print("✅ CUSTOM TEMPLATE DEMO COMPLETE!")
    print("="*70 + "\n")


def demo_template_colors():
    """Show color palettes for all templates"""
    
    print("\n" + "="*70)
    print("🎨 TEMPLATE COLOR PALETTES")
    print("="*70 + "\n")
    
    manager = TemplateManager()
    templates = manager.list_templates(include_custom=False)
    
    for template_info in templates:
        template = manager.load_template(template_info['id'])
        if template:
            print(f"📋 {template['name']}")
            print(f"   Primary:   {template['colors']['primary']}")
            print(f"   Secondary: {template['colors']['secondary']}")
            print(f"   Accent:    {template['colors']['accent']}")
            print(f"   Chart Colors: {', '.join(template['colors']['chart_colors'][:3])}...\n")
    
    print("="*70 + "\n")


if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description="Template System Demo")
    parser.add_argument(
        "--mode",
        choices=["all", "custom", "colors", "single"],
        default="all",
        help="Demo mode to run"
    )
    parser.add_argument(
        "--template",
        default="corporate_blue",
        help="Template name for single mode"
    )
    
    args = parser.parse_args()
    
    if args.mode == "all":
        demo_all_templates()
    elif args.mode == "custom":
        demo_custom_template()
    elif args.mode == "colors":
        demo_template_colors()
    elif args.mode == "single":
        excel_file = "examples/Portfolio Allocation Data.xlsx"
        output_file = f"examples/demo_PPT/single_template_demo.pptx"
        create_presentation_with_template(excel_file, output_file, args.template)
    
    print("\n🎉 Demo complete! Check examples/demo_PPT/ for generated files.")
