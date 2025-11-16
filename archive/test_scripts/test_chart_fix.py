"""Test chart fix for AI Pro tier"""
from src.converter.excel_to_ppt_converter import convert_excel_to_ppt

print("Testing chart fix for AI Pro tier...\n")

result = convert_excel_to_ppt(
    'examples/Company_Data/company_bundle.xlsx', 
    'test_chart_fix.pptx', 
    user_tier='ai_pro', 
    user_ppt_count=0
)

print(f"\n{'='*60}")
print(f"✅ Result: {result.get('success')}")
print(f"📊 Slides created: {result.get('slides_created')}")
print(f"🎨 Template: {result.get('template_used')}")
print(f"{'='*60}")

# Check if charts are in the PPT
if result.get('success'):
    from pptx import Presentation
    prs = Presentation('test_chart_fix.pptx')
    print(f"\nChecking slide 2 (first data slide):")
    if len(prs.slides) > 1:
        slide = prs.slides[1]
        print(f"  Shapes: {len(slide.shapes)}")
        for i, shape in enumerate(slide.shapes):
            print(f"    {i+1}. {shape.shape_type} - {shape.name}")
            if hasattr(shape, 'chart'):
                print(f"       ✅ HAS CHART!")
