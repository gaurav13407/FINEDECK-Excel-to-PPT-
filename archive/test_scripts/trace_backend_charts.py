"""
Complete backend chart flow tracer - shows EXACTLY what's happening
"""
import sys
import pandas as pd
from pptx.util import Inches

# Patch the SimpleFinanceChartBuilder to trace calls
original_create_chart = None
original_detect_chart_type = None

def trace_create_chart(self, slide, df, left, top, width, height, title=None):
    """Wrapped version that traces all parameters"""
    print("\n" + "="*80)
    print("🔍 CHART CREATION TRACED!")
    print("="*80)
    print(f"📍 Called from: {title or 'Unknown'}")
    print(f"📊 DataFrame shape: {df.shape}")
    print(f"📝 DataFrame columns: {list(df.columns)}")
    print(f"🔢 DataFrame dtypes:\n{df.dtypes}")
    print(f"\n📋 Sample data (first 3 rows):")
    print(df.head(3).to_string())
    
    # Call original
    result = original_create_chart(self, slide, df, left, top, width, height, title)
    
    print(f"\n✅ Chart created successfully: {result is not None}")
    print("="*80)
    
    return result

def trace_detect_chart_type(self, df):
    """Wrapped version that traces detection logic"""
    print("\n🔍 DETECTING CHART TYPE...")
    print(f"   DataFrame columns: {list(df.columns)}")
    
    col_names_lower = ' '.join([str(col).lower() for col in df.columns])
    print(f"   Column names (lowercase): '{col_names_lower}'")
    
    # Call original
    chart_type = original_detect_chart_type(self, df)
    
    print(f"   ✅ DETECTED TYPE: {chart_type.upper()}")
    
    # Show why
    time_keywords = ['date', 'quarter', 'month', 'year', 'q1', 'q2', 'q3', 'q4', 'ytd', 'period']
    allocation_keywords = ['allocation', 'portfolio', 'sector', 'distribution', 'breakdown', 'composition']
    performance_keywords = ['top', 'rank', 'performance', 'revenue', 'sales', 'profit', 'growth', 'comparison']
    
    time_matches = [kw for kw in time_keywords if kw in col_names_lower]
    alloc_matches = [kw for kw in allocation_keywords if kw in col_names_lower]
    perf_matches = [kw for kw in performance_keywords if kw in col_names_lower]
    
    if time_matches:
        print(f"   ⏰ Time keywords: {time_matches}")
    if alloc_matches:
        print(f"   🥧 Allocation keywords: {alloc_matches}")
    if perf_matches:
        print(f"   📊 Performance keywords: {perf_matches}")
    
    if not (time_matches or alloc_matches or perf_matches):
        print(f"   ⚠️  NO keywords matched - using default (COLUMN)")
    
    return chart_type

# Patch the module
print("🔧 Patching SimpleFinanceChartBuilder to trace calls...")
from src.converter.simple_finance_charts import SimpleFinanceChartBuilder

original_create_chart = SimpleFinanceChartBuilder.create_chart
original_detect_chart_type = SimpleFinanceChartBuilder.detect_chart_type

SimpleFinanceChartBuilder.create_chart = trace_create_chart
SimpleFinanceChartBuilder.detect_chart_type = trace_detect_chart_type

print("✅ Tracing enabled!")

# Now import and run the actual converter
from src.converter.enhanced_professional_builder import EnhancedProfessionalBuilder
from pptx import Presentation

if len(sys.argv) < 2:
    print("\n❌ Usage: python trace_backend_charts.py <path_to_excel>")
    print("\nExample:")
    print('  python trace_backend_charts.py "examples/Portfolio Allocation Data.xlsx"')
    sys.exit(1)

excel_file = sys.argv[1]

print("\n" + "="*80)
print(f"📂 TESTING WITH: {excel_file}")
print("="*80)

try:
    # Read Excel data
    print("\n📖 Reading Excel file...")
    sheets_data = pd.read_excel(excel_file, sheet_name=None)
    
    # Get first sheet
    first_sheet_name = list(sheets_data.keys())[0]
    data = sheets_data[first_sheet_name]
    
    print(f"✅ Loaded sheet: '{first_sheet_name}'")
    print(f"   Shape: {data.shape}")
    print(f"   Columns: {list(data.columns)}")
    
    # Create converter
    print("\n🏗️ Creating EnhancedProfessionalBuilder...")
    converter = EnhancedProfessionalBuilder()
    
    # Create presentation
    prs = Presentation()
    
    # Prepare template
    template = {"name": "corporate_blue"}
    
    # Call build_presentation which initializes everything
    print("\n🎨 Building presentation...")
    converter.build_presentation(prs, [(first_sheet_name, data)], "Test Project", template)
    
    print("\n" + "="*80)
    print("✅ BUILD COMPLETED!")
    print("="*80)
    
    # Save output
    output_file = "trace_output.pptx"
    prs.save(output_file)
    print(f"\n💾 Saved to: {output_file}")
    
    # Analyze the output
    print("\n📊 ANALYZING OUTPUT...")
    chart_count = 0
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_chart:
                chart_count += 1
                chart = shape.chart
                print(f"\n   Slide {i}: {chart.chart_type} chart found")
                print(f"      Position: left={shape.left} EMUs ({shape.left/914400:.2f} inches)")
                print(f"      Size: {shape.width}x{shape.height} EMUs")
    
    print(f"\n📈 Total charts created: {chart_count}")
    
    if chart_count == 0:
        print("\n❌ NO CHARTS CREATED - Check errors above!")
    else:
        print(f"\n✅ SUCCESS - {chart_count} charts created!")

except FileNotFoundError:
    print(f"\n❌ ERROR: File not found: {excel_file}")
except Exception as e:
    print(f"\n❌ ERROR: {e}")
    import traceback
    traceback.print_exc()
