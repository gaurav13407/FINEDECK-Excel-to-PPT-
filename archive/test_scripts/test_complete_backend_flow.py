"""
COMPLETE END-TO-END TEST
Simulates EXACTLY what the backend does when you upload an Excel file
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pandas as pd
from pptx import Presentation
from pptx.util import Inches

# Step 1: Import the converter (same as backend)
print("\n" + "="*60)
print("STEP 1: Importing Excel to PPT Converter")
print("="*60)
from src.converter.excel_to_ppt_converter import ExcelToPPTConverter

# Step 2: Create test Excel data (simulating user upload)
print("\n" + "="*60)
print("STEP 2: Creating Test Excel Data")
print("="*60)

# Create sample data that should trigger finance charts
test_data = pd.DataFrame({
    'Sector': ['Technology', 'Healthcare', 'Finance', 'Energy', 'Consumer', 'Industrial'],
    'Allocation': [35.5, 25.2, 20.1, 12.3, 4.9, 2.0],
    'Performance': [15.2, 12.8, 10.5, 9.3, 8.1, 6.5]
})

# Save to Excel
excel_path = 'test_backend_flow.xlsx'
test_data.to_excel(excel_path, index=False, sheet_name='Portfolio Data')
print(f"✅ Created test Excel: {excel_path}")
print(f"   Columns: {list(test_data.columns)}")
print(f"   Should detect: 'Sector' + 'Allocation' → PIE chart")

# Step 3: Create converter (same as backend does)
print("\n" + "="*60)
print("STEP 3: Creating ExcelToPPTConverter")
print("="*60)

converter = ExcelToPPTConverter(
    user_tier='ai_pro',  # Use AI_PRO tier to trigger EnhancedProfessionalBuilder
    user_id='test_user',
    user_metadata={
        'name': 'Test User',
        'email': 'test@example.com',
        'company': 'Test Company'
    },
    use_finance_charts=True  # Enable finance charts
)
print(f"✅ Converter created")
print(f"   Tier: ai_pro")
print(f"   Finance charts: ENABLED")

# Step 4: Convert (same as backend calls convert_professional)
print("\n" + "="*60)
print("STEP 4: Converting Excel to PPT")
print("="*60)

output_path = 'test_backend_flow_output.pptx'

print("\n🎯 Calling converter.convert_professional()...")
print("   (This is EXACTLY what your backend does)")
print("")

result = converter.convert_professional(
    excel_path=excel_path,
    output_path=output_path,
    template_name='royal_purple',  # Use your template
    presentation_title='Backend Flow Test',
    user_ppt_count=0,
    use_professional_structure=True  # Force EnhancedProfessionalBuilder
)

print("\n" + "="*60)
print("STEP 5: Checking Results")
print("="*60)

if result['success']:
    print("✅ Conversion SUCCESSFUL!")
    print(f"   Output: {output_path}")
    
    # Analyze the PPT
    print("\n📊 Analyzing generated PPT...")
    prs = Presentation(output_path)
    
    chart_count = 0
    for slide_num, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if hasattr(shape, 'chart'):
                chart_count += 1
                print(f"\n   Slide {slide_num}: Chart found!")
                print(f"      Chart type: {shape.chart.chart_type}")
                print(f"      Position: left={shape.left} EMUs ({shape.left/914400:.2f} inches)")
                print(f"      Position: top={shape.top} EMUs ({shape.top/914400:.2f} inches)")
                
                # Check if position is correct (in millions, not trillions)
                if shape.left < 100000000:  # Less than 100 million EMUs
                    print(f"      ✅ Position CORRECT (in millions of EMUs)")
                else:
                    print(f"      ❌ Position WRONG (in billions/trillions - double Inches bug!)")
    
    print(f"\n📈 Total charts found: {chart_count}")
    
    if chart_count > 0:
        print("\n✅ SUCCESS! Charts were created!")
        print("   Open test_backend_flow_output.pptx to verify they're visible")
    else:
        print("\n❌ WARNING: No charts found in PPT!")
        print("   This means charts are not being created at all")
    
else:
    print("❌ Conversion FAILED!")
    print(f"   Error: {result.get('error', 'Unknown error')}")

# Cleanup
import os
try:
    os.unlink(excel_path)
    print(f"\n🧹 Cleaned up test Excel file")
except:
    pass

print("\n" + "="*60)
print("TEST COMPLETE")
print("="*60)
print("\n📝 Next steps:")
print("   1. Open test_backend_flow_output.pptx")
print("   2. Check if charts are visible on slides")
print("   3. If visible → Backend integration working!")
print("   4. If not visible → Check console output above for errors")
