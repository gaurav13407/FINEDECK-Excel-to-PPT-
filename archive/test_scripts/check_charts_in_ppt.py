"""
Check if charts exist in the downloaded PPT
"""

import sys
import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Get the most recent PPT file from Downloads
downloads_folder = os.path.expanduser("~\\Downloads")

# Find AMZN or TSLA files
ppt_files = []
for file in os.listdir(downloads_folder):
    if file.endswith('.pptx') and ('AMZN' in file or 'TSLA' in file or 'GOOGL' in file):
        full_path = os.path.join(downloads_folder, file)
        ppt_files.append((full_path, os.path.getmtime(full_path)))

if not ppt_files:
    print("❌ No AMZN/TSLA/GOOGL PPT files found in Downloads")
    sys.exit(1)

# Get most recent file
ppt_files.sort(key=lambda x: x[1], reverse=True)
ppt_path = ppt_files[0][0]

print(f"📂 Checking: {os.path.basename(ppt_path)}")
print(f"   Last modified: {ppt_files[0][1]}")
print("=" * 80)

# Open presentation
prs = Presentation(ppt_path)

print(f"\n📊 PRESENTATION ANALYSIS")
print(f"Total slides: {len(prs.slides)}")
print("=" * 80)

# Check each slide
for i, slide in enumerate(prs.slides, 1):
    print(f"\nSlide {i}:")
    
    # Get slide title if exists
    title = "No title"
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text:
            title = shape.text_frame.text
            break
    
    print(f"  Title: {title[:50]}")
    
    # Count shapes by type
    chart_count = 0
    table_count = 0
    text_count = 0
    other_count = 0
    
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.CHART:
            chart_count += 1
            # Get chart details
            try:
                chart = shape.chart
                chart_type = chart.chart_type
                print(f"    ✅ CHART FOUND!")
                print(f"       Type: {chart_type} ({chart_type.value})")
                print(f"       Position: ({shape.left}, {shape.top})")
                print(f"       Size: {shape.width} x {shape.height}")
                print(f"       Series count: {len(chart.series)}")
                if chart.series:
                    print(f"       Series names: {[s.name for s in chart.series]}")
            except Exception as e:
                print(f"    ⚠️ Chart found but error reading: {e}")
        elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            table_count += 1
        elif shape.has_text_frame:
            text_count += 1
        else:
            other_count += 1
    
    print(f"  Shapes: {chart_count} charts, {table_count} tables, {text_count} text, {other_count} other")

print("\n" + "=" * 80)
print("SUMMARY")
print("=" * 80)

total_charts = sum(1 for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART)
print(f"Total charts in presentation: {total_charts}")

if total_charts == 0:
    print("\n❌ NO CHARTS FOUND IN PPT!")
    print("\nPossible causes:")
    print("1. Charts are being created but not saved to the PPT file")
    print("2. Chart objects are None and not being added")
    print("3. The PPT file you're checking is from an older conversion")
    print("\nRecommendation:")
    print("- Delete all PPT files in Downloads")
    print("- Convert a new file from the web app")
    print("- Run this script again")
else:
    print(f"\n✅ Found {total_charts} charts in the PPT!")
    print("\nIf you can't see them:")
    print("1. Charts might be positioned off-screen (check position coordinates above)")
    print("2. Charts might be behind other shapes")
    print("3. Chart size might be too small (check size above)")
