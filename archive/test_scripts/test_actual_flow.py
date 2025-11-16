"""
Test the ACTUAL flow that happens when user uploads through browser
This simulates exactly what happens in production
"""

import sys
import os

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from src.converter.excel_to_ppt_converter import ExcelToPPTConverter

def test_actual_browser_flow():
    """
    Simulate exact browser upload flow:
    1. User uploads Portfolio Allocation Data.xlsx
    2. Selects Royal Purple template
    3. User tier is 'ai_pro' (or 'pro')
    4. Backend calls convert_professional()
    """
    
    print("\n" + "="*70)
    print("🌐 SIMULATING ACTUAL BROWSER UPLOAD FLOW")
    print("="*70)
    
    # EXACT parameters from browser upload
    excel_path = r"C:\Users\gaura\OneDrive\Desktop\Big Projects\FinDeck(Excel to PPT Project)\examples\Portfolio Allocation Data.xlsx"
    output_path = r"C:\Users\gaura\OneDrive\Desktop\Big Projects\FinDeck(Excel to PPT Project)\test_output\browser_simulation.pptx"
    template_name = "royal_purple"  # User selects this in browser
    user_tier = "ai_pro"  # Or "pro"
    presentation_title = "Portfolio Allocation Data"
    
    print(f"\n📊 Input Excel: {os.path.basename(excel_path)}")
    print(f"🎨 Template: {template_name}")
    print(f"👤 User Tier: {user_tier}")
    print(f"📝 Title: {presentation_title}")
    
    # Create converter EXACTLY like backend does
    converter = ExcelToPPTConverter(
        user_tier=user_tier,
        user_id="test_user_123",
        user_metadata={
            'name': 'Test User',
            'email': 'test@findeck.com',
            'company': 'FinDeck Test'
        },
        use_finance_charts=False  # Default value
    )
    
    print("\n🔧 Calling convert_professional() (EXACT backend call)...")
    print("-" * 70)
    
    # Call EXACT method backend uses
    result = converter.convert_professional(
        excel_path=excel_path,
        output_path=output_path,
        template_name=template_name,
        presentation_title=presentation_title,
        user_ppt_count=1,  # User's first PPT this month
        use_professional_structure=True
    )
    
    print("\n" + "="*70)
    print("📊 CONVERSION RESULTS")
    print("="*70)
    
    if result['success']:
        print(f"✅ Success: {result['success']}")
        print(f"📂 Output: {output_path}")
        print(f"📄 Slides Created: {result.get('slides_created', 'N/A')}")
        print(f"📝 Slide Names: {', '.join(result.get('slide_names', []))}")
        
        # Now verify the charts in the PPT
        print("\n" + "="*70)
        print("🔍 VERIFYING CHARTS IN GENERATED PPT")
        print("="*70)
        
        from pptx import Presentation
        prs = Presentation(output_path)
        
        chart_count = 0
        for i, slide in enumerate(prs.slides, 1):
            slide_charts = []
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = shape.chart
                    chart_type_name = str(chart.chart_type).split('.')[-1].replace('(', '').replace(')', '').split()[0]
                    slide_charts.append(f"{chart_type_name} ({chart.chart_type})")
                    chart_count += 1
            
            if slide_charts:
                print(f"   Slide {i}: {len(slide_charts)} chart(s) - {', '.join(slide_charts)}")
        
        if chart_count == 0:
            print("\n❌ NO CHARTS FOUND!")
            print("\n🔍 Debugging:")
            print("   1. Enhanced chart builder initialized?")
            print("   2. Charts being created but returning None?")
            print("   3. Falling back to old method?")
            print("\n💡 Check the logs above for:")
            print("   - '✨ Using EnhancedChartBuilder...'")
            print("   - '📊 Detected chart type: ...'")
            print("   - '✓ Enhanced chart created successfully!'")
        else:
            print(f"\n✅ FOUND {chart_count} CHARTS!")
            
            # Check if they're the NEW enhanced charts
            print("\n🎨 Checking if charts are ENHANCED (purple colors)...")
            from pptx.util import Pt
            from pptx.dml.color import RGBColor
            
            for slide_idx, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.has_chart:
                        chart = shape.chart
                        try:
                            # Try to get first series color
                            if chart.plots[0].series:
                                series = chart.plots[0].series[0]
                                if hasattr(series, 'format') and hasattr(series.format, 'fill'):
                                    fill = series.format.fill
                                    if hasattr(fill, 'fore_color') and hasattr(fill.fore_color, 'rgb'):
                                        rgb = fill.fore_color.rgb
                                        print(f"   Slide {slide_idx} Chart: RGB{rgb} (Purple = CE93D8, 4A148C, 7B1FA2)")
                        except:
                            pass
    else:
        print(f"❌ Conversion Failed!")
        print(f"Error: {result.get('error', 'Unknown error')}")

if __name__ == "__main__":
    test_actual_browser_flow()
