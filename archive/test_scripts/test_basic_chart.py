"""Test chart creation with fallback"""
from src.converter.excel_to_ppt_converter import convert_excel_to_ppt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

print("Testing BASIC tier with Sample_pnl.xlsx...\n")

result = convert_excel_to_ppt(
    'examples/Sample_pnl.xlsx', 
    'test_basic_chart.pptx', 
    user_tier='basic', 
    user_ppt_count=0
)

print(f"\n{'='*60}")
print(f"✅ Success: {result.get('success')}")
print(f"📊 Slides created: {result.get('slides_created')}")
print(f"🎨 Template: {result.get('template_used')}")
print(f"{'='*60}")

# Check if charts are present
if result.get('success'):
    prs = Presentation('test_basic_chart.pptx')
    print(f"\nChecking slides for charts:")
    
    chart_count = 0
    for i, slide in enumerate(prs.slides):
        has_chart = any(shape.shape_type == MSO_SHAPE_TYPE.CHART for shape in slide.shapes)
        if has_chart:
            chart_count += 1
            print(f"  Slide {i+1}: ✅ HAS CHART")
        else:
            print(f"  Slide {i+1}: ❌ No chart")
    
    print(f"\n📊 Total slides with charts: {chart_count}/{len(prs.slides)}")
