"""
Check the most recently downloaded PPT from browser
"""

import os
from pathlib import Path
from pptx import Presentation
from datetime import datetime

# Check common download locations
download_dirs = [
    Path.home() / "Downloads",
    Path(r"C:\Users\gaura\Downloads"),
    Path(r"C:\Users\gaura\OneDrive\Desktop"),
]

print("\n" + "="*80)
print("🔍 SEARCHING FOR RECENT CONVERTED PPT FILES")
print("="*80)

# Find most recent .pptx files
recent_files = []
for download_dir in download_dirs:
    if download_dir.exists():
        print(f"\n📁 Checking: {download_dir}")
        for pptx in download_dir.glob("*.pptx"):
            if "converted" in pptx.name.lower() or any(stock in pptx.name for stock in ["AMZN", "TSLA", "GOOGL"]):
                mod_time = pptx.stat().st_mtime
                recent_files.append((pptx, mod_time))
                print(f"   Found: {pptx.name} ({datetime.fromtimestamp(mod_time).strftime('%H:%M:%S')})")

# Sort by modification time
recent_files.sort(key=lambda x: x[1], reverse=True)

if not recent_files:
    print("\n❌ No converted PPT files found!")
    print("\n💡 Please provide the full path to your downloaded PPT:")
    ppt_path = input("PPT Path: ").strip('"').strip("'")
else:
    # Check the most recent 3 files
    print(f"\n" + "="*80)
    print("📊 ANALYZING MOST RECENT FILES")
    print("="*80)
    
    for pptx_path, mod_time in recent_files[:3]:
        print(f"\n{'='*80}")
        print(f"📄 FILE: {pptx_path.name}")
        print(f"⏰ Modified: {datetime.fromtimestamp(mod_time).strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"📂 Path: {pptx_path}")
        print(f"{'='*80}")
        
        try:
            prs = Presentation(str(pptx_path))
            
            total_slides = len(prs.slides)
            total_charts = 0
            
            print(f"\n📊 Total Slides: {total_slides}")
            
            for slide_idx, slide in enumerate(prs.slides, 1):
                # Get slide title
                slide_title = "Untitled"
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        text = shape.text_frame.text.strip()
                        if len(text) < 100:
                            slide_title = text
                            break
                
                # Check for charts
                chart_count = sum(1 for shape in slide.shapes if shape.has_chart)
                
                if chart_count > 0:
                    total_charts += chart_count
                    print(f"\n   Slide {slide_idx}: {slide_title[:60]}")
                    
                    for shape in slide.shapes:
                        if shape.has_chart:
                            chart = shape.chart
                            chart_type = str(chart.chart_type).split('.')[-1].split('(')[0]
                            
                            # Determine chart category
                            if "PIE" in chart_type:
                                chart_icon = "🥧"
                            elif "BAR" in chart_type:
                                chart_icon = "📊"
                            elif "LINE" in chart_type:
                                chart_icon = "📈"
                            elif "COLUMN" in chart_type:
                                chart_icon = "📊"
                            else:
                                chart_icon = "📉"
                            
                            print(f"      {chart_icon} {chart_type}")
            
            if total_charts == 0:
                print(f"\n❌ NO CHARTS FOUND IN THIS FILE!")
                print(f"\n🔍 Slide breakdown:")
                for slide_idx, slide in enumerate(prs.slides, 1):
                    shape_types = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            shape_types.append("Text")
                        elif shape.has_table:
                            shape_types.append("Table")
                        elif shape.has_chart:
                            shape_types.append("Chart")
                        else:
                            shape_types.append("Other")
                    
                    print(f"   Slide {slide_idx}: {', '.join(set(shape_types))}")
            else:
                print(f"\n✅ FOUND {total_charts} CHART(S) IN {total_slides} SLIDES!")
                
        except Exception as e:
            print(f"\n❌ Error analyzing file: {e}")
            import traceback
            traceback.print_exc()

print("\n" + "="*80)
print("💡 DIAGNOSIS")
print("="*80)

if not recent_files:
    print("\nCouldn't find any converted PPT files.")
    print("Please check your Downloads folder manually.")
else:
    print("\n1. If NO CHARTS were found:")
    print("   → The backend conversion is NOT creating charts")
    print("   → Need to check backend logs for errors")
    print("   → Data might not be suitable for charts")
    print("\n2. If CHARTS were found:")
    print("   → The conversion IS working!")
    print("   → Check if you're opening the correct file")
    print("   → Make sure to open in PowerPoint Desktop, not online")

print("\n" + "="*80)
