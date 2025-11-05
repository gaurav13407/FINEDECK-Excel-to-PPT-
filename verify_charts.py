"""
Verify Charts in Generated Professional PPTs
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

def check_ppt_charts(ppt_path):
    """Check if PPT has charts"""
    if not os.path.exists(ppt_path):
        return None
    
    prs = Presentation(ppt_path)
    
    results = {
        'total_slides': len(prs.slides),
        'slides_with_charts': 0,
        'total_charts': 0,
        'chart_details': []
    }
    
    for i, slide in enumerate(prs.slides, 1):
        slide_charts = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                slide_charts += 1
                results['total_charts'] += 1
                results['chart_details'].append({
                    'slide': i,
                    'chart_title': shape.chart.chart_title.text_frame.text if shape.chart.has_title else 'Untitled'
                })
        
        if slide_charts > 0:
            results['slides_with_charts'] += 1
    
    return results

# Check all generated PPTs
ppts_to_check = [
    'examples/professional_demo/tech_stocks_free_tier.pptx',
    'examples/professional_demo/tech_stocks_basic_tier.pptx',
    'examples/professional_demo/tech_stocks_ai_pro_tier.pptx'
]

print("=" * 80)
print("📊 CHART VERIFICATION - Professional Presentations")
print("=" * 80)

for ppt_path in ppts_to_check:
    tier = ppt_path.split('_')[-2].upper()
    print(f"\n{'='*80}")
    print(f"📄 {tier} TIER: {os.path.basename(ppt_path)}")
    print(f"{'='*80}")
    
    results = check_ppt_charts(ppt_path)
    
    if results is None:
        print("❌ File not found")
        continue
    
    print(f"📊 Total Slides: {results['total_slides']}")
    print(f"📈 Total Charts: {results['total_charts']}")
    print(f"🎨 Slides with Charts: {results['slides_with_charts']}/{results['total_slides']}")
    
    if results['total_charts'] > 0:
        print(f"\n✅ CHARTS FOUND:")
        for detail in results['chart_details']:
            print(f"   Slide {detail['slide']}: {detail['chart_title']}")
    else:
        print("\n⚠️  NO CHARTS FOUND (placeholders only)")
    
    # Check file size
    file_size = os.path.getsize(ppt_path) / 1024
    print(f"\n📦 File Size: {file_size:.1f} KB")
    
    # Verdict
    if results['total_charts'] >= 2:
        print("✅ VERDICT: Good - Multiple charts present")
    elif results['total_charts'] == 1:
        print("⚠️  VERDICT: Acceptable - At least one chart present")
    else:
        print("❌ VERDICT: Needs improvement - No charts found")

print(f"\n{'='*80}")
print("✨ Chart Verification Complete")
print(f"{'='*80}")
